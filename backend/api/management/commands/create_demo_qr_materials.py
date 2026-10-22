import os
import tempfile
import qrcode
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from django.core.management.base import BaseCommand
from django.conf import settings
from django.utils import timezone
from datetime import timedelta
from api.models import Department, Officer, Material, MaterialDocument, AppealStep

YASHNOBOD_CONFIG = [
    {
        "id": "MAT-2026-0001",
        "doc_file": "Ekspertiza.docx",
        "doc_title": "Sud-tibbiy ekspertiza tayinlash qarori (QR-kodli)",
        "citizen_name": "Tillayeva Zaxro Djalalovna",
        "citizen_phone": "+998 90 911-22-33",
        "title_uz": "Tillayeva Zaxro Djalalovnaning tan jarohati yuzasidan sud-tibbiy ekspertiza tayinlash to‘g‘risida",
        "title_ru": "О назначении судебно-медицинской экспертизы по факту телесных повреждений Тиллаевой Захро Джалаловны",
        "officer_id": "off_asror_xazratov",
        "status": "изучаемый",
        "difficulty": 2,
        "days_ago": 4,
        "days_left": 6,
        "steps": [
            ("102 qisqa raqami orqali murojaat qabul qilindi va E-material tizimida ro‘yxatga olindi", "изучаемый"),
            ("Tergovga qadar tekshiruv bo‘linmasi katta inspektori podpolkovnik A.B.Xazratovga biriktirildi", "изучаемый"),
            ("Holat yuzasidan sud-tibbiy ekspertiza tayinlandi va QR-kod bilan elektron tasdiqlandi", "изучаемый")
        ]
    },
    {
        "id": "MAT-2026-0002",
        "doc_file": "Uzaytirish.docx",
        "doc_title": "Muddatni uzaytirish haqida qaror (QR-kodli)",
        "citizen_name": "Musirmonov Islom Bahodir o‘g‘li",
        "citizen_phone": "+998 93 522-44-66",
        "title_uz": "Musirmonov Islom Bahodir o‘g‘lining tan jarohati yuzasidan tekshiruv muddatini uzaytirish to‘g‘risida",
        "title_ru": "О продлении срока доследственной проверки по факту телесных повреждений Мусирмонова Ислома Баходир угли",
        "officer_id": "off_azizjon_umarjonov",
        "status": "изучаемый",
        "difficulty": 3,
        "days_ago": 7,
        "days_left": 14,
        "steps": [
            ("Shoshilinch tibbiy yordam xodimlari orqali 102 raqamiga kelib tushgan xabar ro‘yxatga olindi", "изучаемый"),
            ("Tergovga qadar tekshiruv katta inspektori mayor A.Q.Umarjonovga ijroga yuborildi", "изучаемый"),
            ("Ekspertiza xulosasini kutish munosabati bilan tekshiruv muddati uzaytirildi va prokurorga yuborildi", "изучаемый")
        ]
    },
    {
        "id": "MAT-2026-0003",
        "doc_file": "Otkaz.docx",
        "doc_title": "Jinoyat ishini qo‘zg‘atishni rad etish to‘g‘risida qaror (QR-kodli)",
        "citizen_name": "Mamataliyeva Mohina Zohidovna",
        "citizen_phone": "+998 97 733-55-77",
        "title_uz": "Mamataliyeva Mohina Zohidovnaning arizasi yuzasidan jinoyat ishini qo‘zg‘atishni rad etish to‘g‘risida",
        "title_ru": "Об отказе в возбуждении уголовного дела по заявлению Маматалиевой Мохины Зохидовны",
        "officer_id": "off_asliddin_qodirov",
        "status": "закрыт_в_срок",
        "difficulty": 2,
        "days_ago": 8,
        "days_left": 0,
        "steps": [
            ("Aholidan onlayn mikroqarz masalasidagi murojaat E-material tizimiga kiritildi", "изучаемый"),
            ("Katta inspektor kapitan A.X.Qodirov tomonidan tekshiruv hujjatlari o‘rganib chiqildi", "изучаемый"),
            ("Jinoyat tarkibi bo‘lmaganligi sababli JPK 83-moddasi 2-bandiga asosan jinoyat ishi qo‘zg‘atish rad etildi", "закрыт_в_срок")
        ]
    }
]

OLMAZOR_CONFIG = [
    {
        "id": "MAT-2026-0001",
        "doc_file": "Ekspertiza.docx",
        "doc_title": "Sud-tibbiy ekspertiza tayinlash qarori (QR-kodli)",
        "citizen_name": "Toshmatov Odil Baxtiyorovich",
        "citizen_phone": "+998 90 123-45-67",
        "title_uz": "Toshmatov Odil Baxtiyorovichning tan jarohati yuzasidan sud-tibbiy ekspertiza tayinlash to‘g‘risida",
        "title_ru": "О назначении судебно-медицинской экспертизы по факту телесных повреждений Тошматова Одила Бахтиёровича",
        "officer_id": "off_anvar_egamberdiyev",
        "status": "изучаемый",
        "difficulty": 2,
        "days_ago": 4,
        "days_left": 6,
        "steps": [
            ("102 qisqa raqami orqali murojaat qabul qilindi va E-material tizimida ro‘yxatga olindi", "изучаемый"),
            ("Tergovga qadar tekshiruv bo‘linmasi katta inspektori mayor A.E.Egamberdiyevga biriktirildi", "изучаемый"),
            ("Holat yuzasidan sud-tibbiy ekspertiza tayinlandi va QR-kod bilan elektron tasdiqlandi", "изучаемый")
        ]
    },
    {
        "id": "MAT-2026-0002",
        "doc_file": "Uzaytirish.docx",
        "doc_title": "Muddatni uzaytirish haqida qaror (QR-kodli)",
        "citizen_name": "Salimov Timur Azizovich",
        "citizen_phone": "+998 91 234-56-78",
        "title_uz": "Salimov Timur Azizovichning murojaati yuzasidan tekshiruv muddatini uzaytirish to‘g‘risida",
        "title_ru": "О продлении срока доследственной проверки по факту Салимова Тимура Азизовича",
        "officer_id": "off_jovohir_muydinov",
        "status": "изучаемый",
        "difficulty": 3,
        "days_ago": 7,
        "days_left": 14,
        "steps": [
            ("Shoshilinch tibbiy yordam xodimlari orqali 102 raqamiga kelib tushgan xabar ro‘yxatga olindi", "изучаемый"),
            ("Tergovga qadar tekshiruv katta inspektori mayor J.Y.Muydinovga ijroga yuborildi", "изучаемый"),
            ("Ekspertiza xulosasini kutish munosabati bilan tekshiruv muddati uzaytirildi va prokurorga yuborildi", "изучаемый")
        ]
    },
    {
        "id": "MAT-2026-0003",
        "doc_file": "Otkaz.docx",
        "doc_title": "Jinoyat ishini qo‘zg‘atishni rad etish to‘g‘risida qaror (QR-kodli)",
        "citizen_name": "Ismoilov Sardor Rustamovich",
        "citizen_phone": "+998 93 345-67-89",
        "title_uz": "Ismoilov Sardor Rustamovichning arizasi yuzasidan jinoyat ishini qo‘zg‘atishni rad etish to‘g‘risida",
        "title_ru": "Об отказе в возбуждении уголовного дела по заявлению Исмоилова Сардора Рустамовича",
        "officer_id": "off_shaxzod_toxirov",
        "status": "закрыт_в_срок",
        "difficulty": 2,
        "days_ago": 8,
        "days_left": 0,
        "steps": [
            ("Aholidan onlayn mikroqarz masalasidagi murojaat E-material tizimiga kiritildi", "изучаемый"),
            ("Inspektor kapitan Sh.B.Toxirov tomonidan tekshiruv hujjatlari o‘rganib chiqildi", "изучаемый"),
            ("Jinoyat tarkibi bo‘lmaganligi sababli JPK 83-moddasi 2-bandiga asosan jinoyat ishi qo‘zg‘atish rad etildi", "закрыт_в_срок")
        ]
    }
]

def stamp_document_with_qr(source_docx_path, dest_docx_path, qr_image_path, material_id, decision_title, officer_name, district='yashnobod'):
    doc = docx.Document(source_docx_path)

    # Add stamp table
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    cell_text = table.rows[0].cells[0]
    cell_qr = table.rows[0].cells[1]
    cell_text.width = Inches(4.3)
    cell_qr.width = Inches(1.8)

    dept_label = "Yashnobod tumani bo'yicha IIO FMB" if district == 'yashnobod' else "Olmazor tumani bo'yicha IIO FMB"

    p1 = cell_text.paragraphs[0]
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run(f"O'ZBEKISTON RESPUBLIKASI ICHKI ISHLAR VAZIRLIGI\n{dept_label} TERGOV BO'LIMI\nTERGOVGA QADAR TEKSHIRUV BO'LINMASI\n")
    r1.bold = True
    r1.font.size = Pt(8)
    r1.font.color.rgb = RGBColor(15, 23, 42)

    p2 = cell_text.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    runs_info = [
        ("Hujjat nomi: ", True), (f"{decision_title}\n", False),
        ("Elektron ro'yxat raqami: ", True), (f"{material_id}\n", False),
        ("Tasdiqlovchi mas'ul xodim: ", True), (f"{officer_name}\n", False),
        ("Tasdiqlash holati: ", True), ("Elektron raqamli imzo va QR-verifikatsiya bilan tasdiqlangan\n", False),
        ("Material havolasi: ", True), (f"https://tqtb-olmazor.uz/{district}/arizalarim?id={material_id}", False),
    ]
    for text, bold in runs_info:
        r = p2.add_run(text)
        r.bold = bold
        r.font.size = Pt(7.5)
        r.font.color.rgb = RGBColor(55, 65, 81)

    p_qr = cell_qr.paragraphs[0]
    p_qr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_qr.paragraph_format.space_after = Pt(2)
    r_qr = p_qr.add_run()
    r_qr.add_picture(qr_image_path, width=Inches(1.3), height=Inches(1.3))

    p_lbl = cell_qr.add_paragraph("QR-KODNI SKANERLANG\n(Materialni ochish)")
    p_lbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_lbl.runs[0].bold = True
    p_lbl.runs[0].font.size = Pt(6.5)
    p_lbl.runs[0].font.color.rgb = RGBColor(100, 116, 139)

    doc.save(dest_docx_path)


class Command(BaseCommand):
    help = 'Creates demo materials with QR-code embedded official documents for Yashnobod or Olmazor'

    def add_arguments(self, parser):
        parser.add_argument('--district', default='yashnobod', choices=['yashnobod', 'olmazor'], help='District to populate (yashnobod or olmazor)')
        parser.add_argument('--database', default=None, help='Database alias to use')

    def handle(self, *args, **options):
        from backend.db_router import set_current_branch
        district = options.get('district', 'yashnobod')
        db = options.get('database') or ('default' if district == 'olmazor' else 'yashnobod')
        set_current_branch(db)

        dept_id = district
        dept_name = "Olmazor tumani" if district == 'olmazor' else "Yashnobod tumani"
        dept_name_ru = "Алмазарский район" if district == 'olmazor' else "Яшнабадский район"
        dept, _ = Department.objects.get_or_create(id=dept_id, defaults={"name_uz": dept_name, "name_ru": dept_name_ru})

        configs = OLMAZOR_CONFIG if district == 'olmazor' else YASHNOBOD_CONFIG

        # Find documents directory
        docs_folder = '/home/ubuntu/e-material/yashnobod_docs'
        if not os.path.exists(docs_folder):
            docs_folder = '/Users/macbook/Downloads'

        media_docs_dir = os.path.join(settings.MEDIA_ROOT, 'material_docs')
        os.makedirs(media_docs_dir, exist_ok=True)

        now = timezone.now()

        for config in configs:
            mat_id = config["id"]
            officer = Officer.objects.filter(id=config["officer_id"]).first()
            if not officer:
                officer = Officer.objects.filter(department=dept, role='investigator').first()

            reg_date = now - timedelta(days=config["days_ago"])
            deadline_date = reg_date + timedelta(days=10 + config["days_left"])

            material, _ = Material.objects.get_or_create(
                id=mat_id,
                defaults={
                    "citizen_name": config["citizen_name"],
                    "citizen_phone": config["citizen_phone"],
                    "title_uz": config["title_uz"],
                    "title_ru": config["title_ru"],
                    "registered_at": reg_date,
                    "deadline": deadline_date,
                    "status": config["status"],
                    "difficulty": config["difficulty"],
                    "department": dept,
                    "officer": officer,
                    "material_type": "e_material",
                    "source_from": "102",
                    "closed_at": now if config["status"] == "закрыт_в_срок" else None,
                }
            )
            # Ensure properties match
            material.citizen_name = config["citizen_name"]
            material.citizen_phone = config["citizen_phone"]
            material.title_uz = config["title_uz"]
            material.title_ru = config["title_ru"]
            material.status = config["status"]
            material.difficulty = config["difficulty"]
            material.department = dept
            material.officer = officer
            material.registered_at = reg_date
            material.deadline = deadline_date
            if config["status"] == "закрыт_в_срок":
                material.closed_at = now
            material.save()

            # Create AppealSteps
            material.appeals.all().delete()
            for idx, (step_title, step_status) in enumerate(config["steps"]):
                AppealStep.objects.create(
                    material=material,
                    status=step_title,
                    time=reg_date + timedelta(days=idx)
                )

            # Generate QR Code image
            qr_url = f"https://tqtb-olmazor.uz/{district}/arizalarim?id={mat_id}"
            qr = qrcode.QRCode(box_size=5, border=1)
            qr.add_data(qr_url)
            qr.make(fit=True)
            qr_img = qr.make_image(fill_color="black", back_color="white")

            with tempfile.NamedTemporaryFile(suffix='.png', delete=False) as tmp_qr:
                qr_img.save(tmp_qr.name)
                tmp_qr_path = tmp_qr.name

            # Source doc path
            src_doc = os.path.join(docs_folder, config["doc_file"])
            dest_filename = f"{district}_{mat_id}_{config['doc_file']}"
            dest_doc_path = os.path.join(media_docs_dir, dest_filename)

            officer_name = f"{officer.rank_uz} {officer.name_uz}" if officer else "Tergovchi"
            try:
                if os.path.exists(src_doc):
                    stamp_document_with_qr(
                        src_doc,
                        dest_doc_path,
                        tmp_qr_path,
                        mat_id,
                        config["doc_title"],
                        officer_name,
                        district=district
                    )
                else:
                    # Create blank docx with content if source missing
                    doc = docx.Document()
                    doc.add_heading(config["doc_title"], 0)
                    doc.add_paragraph(config["title_uz"])
                    doc.save(dest_doc_path)
                    stamp_document_with_qr(
                        dest_doc_path,
                        dest_doc_path,
                        tmp_qr_path,
                        mat_id,
                        config["doc_title"],
                        officer_name,
                        district=district
                    )

                # Link MaterialDocument
                relative_path = os.path.join('material_docs', dest_filename)
                MaterialDocument.objects.filter(material=material, original_name=config["doc_title"] + ".docx").delete()
                doc_obj = MaterialDocument.objects.create(
                    material=material,
                    original_name=config["doc_title"] + ".docx",
                    uploaded_by=officer_name,
                )
                doc_obj.file.name = relative_path
                doc_obj.save()

                self.stdout.write(self.style.SUCCESS(f"✓ Created {mat_id} for {district} with QR document: {dest_filename}"))

            finally:
                if os.path.exists(tmp_qr_path):
                    os.remove(tmp_qr_path)

        self.stdout.write(self.style.SUCCESS(f"Successfully populated {district} demo materials with QR codes!"))
